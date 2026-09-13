"""
Core unit tests for the extracted search engine primitives.

These tests do not need a server, a corpus or Flask. They pin the behavior
that the whole product depends on: tokenization, query interpretation,
filename safety, phrase counting, snippet/highlight construction and
pagination.

They also assert parity against the Werkzeug original for the vendored
`secure_filename`, so "vendored" stays a provable claim rather than a
hope.
"""

import os

import pytest

from search_engine.config import (
    SUPPORTED_EXTENSIONS,
)
from search_engine.extraction import (
    extract_pages,
    extract_text,
)
from search_engine.filenames import (
    normalize_requested_filenames,
    sanitize_upload_filename,
    secure_filename,
)
from search_engine.pagination import (
    empty_query_response,
    paginate,
)
from search_engine.query import (
    normalize_search_query,
    parse_filetype_filter,
)
from search_engine.snippets import (
    build_snippet_result,
    count_phrase_occurrences,
    get_snippet_and_page,
)
from search_engine.tokenizer import (
    tokenize,
    tokenize_filename,
)


# ============================================================
# TOKENIZERS
# ============================================================

def test_tokenize_drops_single_characters():

    assert tokenize("A bb C ddd") == ["bb", "ddd"]


def test_tokenize_keeps_numerics_above_one_character():

    assert tokenize("module 12 of 5") == ["module", "12", "of"]


def test_tokenize_strips_punctuation():

    assert tokenize("network-layer, TCP/IP!") == [
        "network",
        "layer",
        "tcp",
        "ip",
    ]


def test_tokenize_filename_keeps_single_characters():

    assert tokenize_filename("Module 2 a") == [
        "module",
        "2",
        "a",
    ]


def test_tokenize_filename_lowercases():

    assert tokenize_filename("BCS502 Notes") == [
        "bcs502",
        "notes",
    ]


# ============================================================
# QUERY NORMALIZATION
# ============================================================

@pytest.mark.parametrize(
    "raw, expected",
    [
        ("", ""),
        ("   ", ""),
        ("network", "network"),
        ("BCS502 Module 2.pdf", "BCS502 Module 2"),
        ("BCS502 Module 2.PDF", "BCS502 Module 2"),
        ('"BCS502 Module 2.pdf"', '"BCS502 Module 2"'),
        ('"network"', '"network"'),
        ("notes.docx", "notes"),
        ("notes.txt", "notes"),
        # Only a trailing extension is removed.
        ("pdf notes", "pdf notes"),
        # An extension in the middle is untouched.
        ("a.pdf b", "a.pdf b"),
    ],
)
def test_normalize_search_query(raw, expected):

    assert normalize_search_query(raw) == expected


# ============================================================
# FILETYPE FILTER
# ============================================================

@pytest.mark.parametrize(
    "raw, expected",
    [
        ("network", ("network", None)),
        ("pdf", ("", "pdf")),
        (".pdf", ("", "pdf")),
        ("network pdf", ("network", "pdf")),
        ("network .pdf", ("network", "pdf")),
        ("BCS502 Module 2.pdf", ("BCS502 Module 2", "pdf")),
        ('"BCS502 Module 2.pdf"', ('"BCS502 Module 2"', "pdf")),
        ("customer .docx", ("customer", "docx")),
        ("notes txt", ("notes", "txt")),
        # First type token wins, later ones are dropped as tokens.
        ("pdf docx notes", ("notes", "pdf")),
        ("", ("", None)),
    ],
)
def test_parse_filetype_filter(raw, expected):

    assert parse_filetype_filter(raw) == expected


def test_supported_extensions_are_pdf_docx_txt():

    assert SUPPORTED_EXTENSIONS == (".pdf", ".docx", ".txt")


# ============================================================
# FILENAME SAFETY
# ============================================================

FILENAME_CORPUS = [
    "simple.txt",
    "My cool movie.mov",
    "../../../etc/passwd",
    "..\\..\\windows\\system32\\config.txt",
    "i contain cool \xfcml\xe4uts.txt",
    "/absolute/path/notes.pdf",
    "nested/dir/report.pdf",
    "...leading-dots.txt",
    "trailing-dots....txt",
    "under_score.docx",
    "dash-es.pdf",
    "spaced name final.pdf",
    "",
    ".",
    "..",
    "\x00null.txt",
    "emoji-\U0001f600-notes.txt",
    "tab\tname.txt",
    "con.txt",
    "nul.pdf",
    "COM1.docx",
    "a" * 200 + ".txt",
    "日本語ファイル.txt",
    "mixed Case File.PDF",
    "quote'\"backtick.txt",
    "percent%20encoded.txt",
    "plus+and=equals.txt",
    "$dollar.txt",
    "hash#tag.txt",
]


@pytest.mark.parametrize("filename", FILENAME_CORPUS)
def test_secure_filename_matches_werkzeug(filename):
    """
    The core vendors secure_filename so it has no web-stack dependency.

    Parity with Werkzeug is asserted, not assumed. Werkzeug remains
    installed because the Flask adapter needs it, so this comparison is
    always available.
    """

    from werkzeug.utils import secure_filename as werkzeug_secure_filename

    assert secure_filename(filename) == werkzeug_secure_filename(
        filename
    )


@pytest.mark.parametrize(
    "filename, expected",
    [
        ("notes.pdf", "notes.pdf"),
        ("../../evil.pdf", "evil.pdf"),
        ("dir/report.docx", "report.docx"),
        ("My Cool Notes.PDF", "My_Cool_Notes.PDF"),
        # Unsupported extension -> reject.
        ("malware.exe", ""),
        ("archive.zip", ""),
        ("noextension", ""),
        ("", ""),
        (None, ""),
        # A name that sanitizes to nothing is rejected.
        ("///", ""),
    ],
)
def test_sanitize_upload_filename(filename, expected):

    assert sanitize_upload_filename(filename) == expected


def test_sanitize_upload_filename_never_keeps_a_separator():

    for filename in FILENAME_CORPUS:

        safe = sanitize_upload_filename(filename)

        assert os.sep not in safe
        assert "/" not in safe
        assert "\\" not in safe


def test_normalize_requested_filenames_dedupes_and_strips():

    assert normalize_requested_filenames([
        " b.txt ",
        "b.txt",
        "a.txt",
        7,
        None,
        "",
        "dir/c.txt",
    ]) == ["b.txt", "a.txt", "c.txt"]


# ============================================================
# PHRASE COUNTING
# ============================================================

@pytest.mark.parametrize(
    "text, phrase, expected",
    [
        ("", "", 0),
        ("network layer", "", 0),
        ("", "network", 0),
        ("network layer network layer", "network layer", 2),
        ("network   layer", "network layer", 1),
        ("NETWORK LAYER", "network layer", 1),
        # Adjacency is required; proximity is not a phrase.
        ("network and layer", "network layer", 0),
        # Non-overlapping counting.
        ("aa aa aa", "aa aa", 1),
    ],
)
def test_count_phrase_occurrences(text, phrase, expected):

    assert count_phrase_occurrences(text, phrase) == expected


# ============================================================
# SNIPPETS
# ============================================================

def page_text_index():
    return {
        "doc.pdf": [
            {
                "page": 1,
                "text": (
                    "introduction to the OSI reference model "
                    "and its seven layers"
                ),
            },
            {
                "page": 2,
                "text": (
                    "routing protocols move packets across a "
                    "computer network using identifiers such as 999"
                ),
            },
        ],
        "empty.pdf": [],
    }


def test_snippet_no_pages_reports_no_readable_text():

    result = get_snippet_and_page(
        page_text_index(),
        "missing.pdf",
        ["network"],
    )

    assert result == {
        "snippet": "No readable text found.",
        "page": None,
        "highlights": [],
    }


def test_snippet_without_query_words_reports_no_match():

    result = get_snippet_and_page(
        page_text_index(),
        "doc.pdf",
        [],
    )

    assert result == {
        "snippet": "No matching text found.",
        "page": None,
        "highlights": [],
    }


def test_snippet_prefers_exact_phrase_and_reports_its_page():

    result = get_snippet_and_page(
        page_text_index(),
        "doc.pdf",
        ["computer", "network"],
    )

    assert result["page"] == 2
    assert "computer network" in result["snippet"]
    assert len(result["highlights"]) == 1

    span = result["highlights"][0]

    assert result["snippet"][span["start"]:span["end"]] == (
        "computer network"
    )


def test_snippet_exact_token_beats_prefix():

    result = get_snippet_and_page(
        page_text_index(),
        "doc.pdf",
        ["routing"],
    )

    assert result["page"] == 2
    assert "routing" in result["snippet"]


def test_snippet_prefix_match_is_highlighted():

    result = get_snippet_and_page(
        page_text_index(),
        "doc.pdf",
        ["rout"],
    )

    assert result["page"] == 2
    assert len(result["highlights"]) == 1

    span = result["highlights"][0]

    assert result["snippet"][
        span["start"]:span["end"]
    ].startswith("rout")


def test_snippet_numeric_token():

    result = get_snippet_and_page(
        page_text_index(),
        "doc.pdf",
        ["999"],
    )

    assert result["page"] == 2
    assert "999" in result["snippet"]


def test_snippet_windows_add_ellipsis():

    long_text = " ".join(
        f"word{index}" for index in range(200)
    ) + " needle " + " ".join(
        f"tail{index}" for index in range(200)
    )

    result = build_snippet_result(
        long_text,
        1,
        long_text.index("needle"),
        len("needle"),
        140,
        "needle",
    )

    assert result["snippet"].startswith("... ")
    assert result["snippet"].endswith(" ...")
    assert len(result["snippet"]) < len(long_text)
    assert result["highlights"]


def test_snippet_highlights_are_merged_when_touching():

    text = "alpha beta gamma delta"

    result = build_snippet_result(
        text,
        3,
        text.index("beta gamma"),
        len("beta gamma"),
        140,
        "beta gamma",
    )

    assert result["page"] == 3
    assert result["highlights"] == [
        {
            "start": text.index("beta gamma"),
            "end": text.index("beta gamma") + len("beta gamma"),
        }
    ]


# ============================================================
# PAGINATION
# ============================================================

def ranked(count):
    return [{"title": f"doc{index}"} for index in range(count)]


def test_paginate_first_page():

    payload = paginate(ranked(5), 1, 2)

    assert payload["pagination"] == {
        "page": 1,
        "limit": 2,
        "total": 5,
        "total_pages": 3,
        "has_next": True,
        "has_previous": False,
        "start": 1,
        "end": 2,
    }

    assert [item["title"] for item in payload["results"]] == [
        "doc0",
        "doc1",
    ]


def test_paginate_middle_page():

    payload = paginate(ranked(5), 2, 2)

    assert payload["pagination"]["has_previous"] is True
    assert payload["pagination"]["has_next"] is True
    assert payload["pagination"]["start"] == 3
    assert payload["pagination"]["end"] == 4


def test_paginate_clamps_beyond_last_page():

    payload = paginate(ranked(5), 99, 2)

    assert payload["pagination"]["page"] == 3
    assert payload["pagination"]["has_next"] is False
    assert [item["title"] for item in payload["results"]] == ["doc4"]


def test_paginate_empty_results():

    payload = paginate([], 4, 10)

    assert payload["results"] == []
    assert payload["pagination"] == {
        "page": 4,
        "limit": 10,
        "total": 0,
        "total_pages": 0,
        "has_next": False,
        "has_previous": False,
        "start": 0,
        "end": 0,
    }


def test_empty_query_response_shape_is_preserved():
    """
    A blank query is answered before ranking, and historically its
    pagination block carries no start/end offsets. Pinned so a refactor
    cannot silently change the public API shape.
    """

    payload = empty_query_response()

    assert payload == {
        "results": [],
        "pagination": {
            "page": 1,
            "limit": 10,
            "total": 0,
            "total_pages": 0,
            "has_next": False,
            "has_previous": False,
        },
    }


# ============================================================
# EXTRACTION
# ============================================================

def test_extract_txt_single_logical_page(tmp_path):

    path = tmp_path / "notes.txt"
    path.write_text(
        "first line\nsecond line\n",
        encoding="utf-8",
    )

    text = extract_text(str(path), "notes.txt")

    assert text == "first line\nsecond line\n"

    pages = extract_pages(str(path), "notes.txt")

    assert len(pages) == 1
    assert pages[0]["page"] == 1
    assert pages[0]["text"] == "first line\nsecond line\n"


def test_extract_empty_txt_has_no_pages(tmp_path):

    path = tmp_path / "blank.txt"
    path.write_text("   \n", encoding="utf-8")

    pages = extract_pages(str(path), "blank.txt")

    assert pages == []


def test_extract_docx_reads_paragraphs(tmp_path):

    import docx

    path = tmp_path / "notes.docx"

    document = docx.Document()
    document.add_paragraph("Computer Network Technologies")
    document.add_paragraph("Routing and switching")
    document.save(str(path))

    text = extract_text(str(path), "notes.docx")

    assert "computer network technologies" in text
    assert "routing and switching" in text

    pages = extract_pages(str(path), "notes.docx")

    assert len(pages) == 1
    assert "computer network technologies" in pages[0]["text"]


def test_extract_unsupported_type_yields_nothing(tmp_path):

    path = tmp_path / "data.csv"
    path.write_text("a,b,c\n", encoding="utf-8")

    assert extract_text(str(path), "data.csv") == ""
    assert extract_pages(str(path), "data.csv") == []
