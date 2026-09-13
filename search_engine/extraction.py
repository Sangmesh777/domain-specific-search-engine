"""
Document text extraction.

Two views of the same document are produced:

    extract_text  -> one lowercased string (content term frequencies)
    extract_pages -> [{"page": n, "text": "..."}] (snippets, page numbers)

PDFs are page-accurate. DOCX and TXT have no fixed page geometry, so they
are treated as one logical page. This is existing product behavior and is
preserved deliberately.

Extraction never raises for an unreadable document: it logs and returns
empty text, and the caller decides whether an empty document is an error
(indexing a document with no readable text is rejected upstream).
"""

import PyPDF2
import docx


def extract_text(file_path, filename):
    """Extract complete text from PDF, DOCX, or TXT."""

    text = ""

    try:
        filename_lower = filename.lower()

        if filename_lower.endswith(".pdf"):

            with open(file_path, "rb") as file:
                reader = PyPDF2.PdfReader(file)

                for page in reader.pages:

                    extracted = page.extract_text()

                    if extracted:
                        text += extracted + " "

        elif filename_lower.endswith(".docx"):

            document = docx.Document(file_path)

            for paragraph in document.paragraphs:
                text += paragraph.text + " "

        elif filename_lower.endswith(".txt"):

            with open(
                file_path,
                "r",
                encoding="utf-8"
            ) as file:
                text = file.read()

    except Exception as error:

        print(
            f"[ERROR] Could not read {filename}: {error}"
        )

    return text.lower()


def extract_pages(file_path, filename):
    """
    Extract page-level text.

    PDFs:
        returns [{"page": 1, "text": "..."}]

    DOCX/TXT:
        treated as one logical page.
    """

    pages = []

    try:

        filename_lower = filename.lower()

        if filename_lower.endswith(".pdf"):

            with open(file_path, "rb") as file:

                reader = PyPDF2.PdfReader(file)

                for page_number, page in enumerate(
                    reader.pages,
                    start=1
                ):

                    extracted = page.extract_text()

                    if extracted:
                        pages.append({
                            "page": page_number,
                            "text": extracted.lower()
                        })

        elif filename_lower.endswith(".docx"):

            document = docx.Document(file_path)

            text = "\n".join(
                paragraph.text
                for paragraph in document.paragraphs
            )

            if text.strip():
                pages.append({
                    "page": 1,
                    "text": text.lower()
                })

        elif filename_lower.endswith(".txt"):

            with open(
                file_path,
                "r",
                encoding="utf-8"
            ) as file:

                text = file.read()

                if text.strip():
                    pages.append({
                        "page": 1,
                        "text": text.lower()
                    })

    except Exception as error:

        print(
            f"[ERROR] Could not extract pages from "
            f"{filename}: {error}"
        )

    return pages
