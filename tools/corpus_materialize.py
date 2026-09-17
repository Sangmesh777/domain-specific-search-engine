"""
Materialise the golden corpus onto disk as real TXT / DOCX / PDF files.

The files are written through the real libraries the engine uses for
extraction (python-docx, PyPDF2-compatible PDF), so uploading them
exercises the true extraction path instead of bypassing it.
"""

from pathlib import Path

from tools.golden_corpus import (
    FORMAT_DOCX,
    FORMAT_PDF,
    FORMAT_TXT,
    GOLDEN_CORPUS,
)


class CorpusMaterialisationError(Exception):
    """Raised when a corpus document cannot be materialised."""


def _write_txt(path, text):
    path.write_text(text, encoding="utf-8")


def _write_docx(path, text):
    import docx

    document = docx.Document()

    for line in text.split("\n"):
        document.add_paragraph(line)

    document.save(str(path))


def _write_pdf(path, text):
    from fpdf import FPDF

    try:
        text.encode("latin-1")
    except UnicodeEncodeError as error:
        raise CorpusMaterialisationError(
            f"PDF document {path.name!r} contains characters outside "
            f"latin-1 that the core PDF font cannot embed: {error}"
        ) from error

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)

    for line in text.split("\n"):
        pdf.cell(0, 8, line)
        pdf.ln(8)

    pdf.output(str(path))


_WRITERS = {
    FORMAT_TXT: _write_txt,
    FORMAT_DOCX: _write_docx,
    FORMAT_PDF: _write_pdf,
}


def materialize_corpus(destination):
    """
    Write every golden-corpus document into `destination`.

    Returns the list of written paths in definition order.
    """

    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)

    written = []

    for document in GOLDEN_CORPUS:
        name = document["name"]
        fmt = document["format"]

        writer = _WRITERS.get(fmt)

        if writer is None:
            raise CorpusMaterialisationError(
                f"Unsupported corpus format {fmt!r} for {name!r}"
            )

        path = destination / name
        writer(path, document["text"])
        written.append(path)

    return written


if __name__ == "__main__":
    import sys

    target = sys.argv[1] if len(sys.argv) > 1 else "build/golden_corpus"

    for written_path in materialize_corpus(target):
        print("wrote", written_path)
